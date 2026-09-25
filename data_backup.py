"""Portable, verified data bundles. Restore always creates a separate directory."""
import hashlib
import json
import os
import re
import tempfile
import zipfile
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

from config import Settings, app_data_dir, map_session_paths
from core import ExcelLedger, LedgerError, image_digest
from image_io import supported_image
from session_store import SessionStore
from storage import folder_files, writable_directory
from version import APP_VERSION


def validate_session(session):
    from batch import QueueEntry, ReceiptDraft
    if (not isinstance(session, dict) or session.get("version") != 1
            or not isinstance(session.get("entries"), list) or not isinstance(session.get("workbook"), str)
            or not isinstance(session.get("view", {}), dict)):
        raise ValueError("バックアップ内の下書きの形式が不正です。")
    def draft_valid(raw):
        if raw is None:
            return
        draft = ReceiptDraft(**raw)
        if (any(not isinstance(value, str) for key, value in asdict(draft).items() if key != "items")
                or not isinstance(draft.items, list) or any(not isinstance(row, list) or len(row) != 5
                or any(not isinstance(value, str) for value in row) for row in draft.items)):
            raise ValueError("バックアップ内の商品明細の形式が不正です。")
    def entries_valid(entries):
        seen = set()
        for raw in entries:
            entry = QueueEntry(**raw)
            if (not isinstance(entry.path, str) or not entry.path or entry.path.casefold() in seen
                    or entry.state not in ("pending", "reading", "ready", "failed", "saved")
                    or any(not isinstance(value, str) for value in (entry.stored_image, entry.image_hash, entry.saved_to))):
                raise ValueError("バックアップ内の画像一覧の形式が不正です。")
            seen.add(entry.path.casefold())
            draft_valid(entry.draft)
    entries_valid(session["entries"])
    draft_valid(session.get("form"))
    batches = session.get("removed_batches", [])
    if not isinstance(batches, list) or len(batches) > 10:
        raise ValueError("バックアップ内の削除履歴の形式が不正です。")
    for batch in batches:
        if not isinstance(batch, dict) or not isinstance(batch.get("entries"), list) or not isinstance(batch.get("view", {}), dict):
            raise ValueError("バックアップ内の削除履歴の形式が不正です。")
        for item in batch["entries"]:
            if not isinstance(item.get("index"), int) or item["index"] < 0:
                raise ValueError("バックアップ内の削除履歴の順序が不正です。")
        entries_valid([item["entry"] for item in batch["entries"]])


def write_json_atomic(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".state-", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(data, handle, ensure_ascii=False)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def recover_activation():
    journal = app_data_dir() / "restore-pending.json"
    if not journal.exists():
        return
    data = json.loads(journal.read_text(encoding="utf-8"))
    Settings(**data["settings"]).save()
    SessionStore(app_data_dir()).save(data["session"])
    journal.unlink()


def activate_restored_state(settings, session):
    write_json_atomic(app_data_dir() / "restore-pending.json", {"settings": settings, "session": session})
    recover_activation()


def session_image_paths(session):
    result = []
    def visit(value):
        if isinstance(value, dict):
            for key, item in value.items():
                if key in ("image", "path", "stored_image") and isinstance(item, str) and item and supported_image(item):
                    result.append(item)
                else:
                    visit(item)
        elif isinstance(value, list):
            for item in value:
                visit(item)
    visit(session)
    return result


def create_bundle(destination, settings, session):
    destination = Path(destination).resolve()
    sources, missing = {}, set()
    def add(source, kind):
        if not source:
            return
        path = Path(source).resolve()
        if path == destination:
            raise ValueError("バックアップ先には元データと異なるZIPファイルを指定してください。")
        key = str(path).casefold()
        if path.is_file():
            sources.setdefault(key, (path, kind))
        else:
            missing.add(str(path))

    workbook = Path(settings.workbook_path).resolve()
    ledgers = {str(workbook)}
    def collect_ledgers(value):
        if isinstance(value, dict):
            for key, item in value.items():
                if key in ("workbook", "saved_to") and isinstance(item, str) and item:
                    ledgers.add(str(Path(item).resolve()))
                else:
                    collect_ledgers(item)
        elif isinstance(value, list):
            for item in value:
                collect_ledgers(item)
    collect_ledgers(session)
    for path in sorted(ledgers):
        if not Path(path).exists():
            if Path(path) != workbook:
                missing.add(path)
            continue
        add(path, "excel")
        for receipt in ExcelLedger(path).load():
            add(receipt.image_path, "image")
        # Include this workbook's automatic/safety copies, not unrelated documents.
        prefix = Path(path).stem
        for backup in Path(path).parent.iterdir():
            if re.fullmatch(re.escape(prefix) + r"\.(backup|before-delete|before-restore)-\d{8}-\d{6}-\d{6}\.xlsx", backup.name):
                add(backup, "backup")
                for receipt in ExcelLedger(backup).load():
                    add(receipt.image_path, "image")
    for path in folder_files(settings.image_folder):
        if supported_image(path):
            add(path, "image")
    for path in session_image_paths(session):
        add(path, "image")
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=".bundle-", suffix=".zip", dir=destination.parent)
    os.close(fd)
    manifest = {"format": "ReceiptLedgerData", "version": 1, "app_version": APP_VERSION,
                "created": datetime.now().isoformat(), "settings": asdict(settings), "session": session,
                "files": [], "missing": sorted(missing)}
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for index, (source, kind) in enumerate(sources.values()):
                name = f"files/{index:06d}{source.suffix.lower()}"
                digest = hashlib.sha256()
                size = 0
                with source.open("rb") as input_file, archive.open(name, "w", force_zip64=True) as output:
                    while chunk := input_file.read(1024 * 1024):
                        digest.update(chunk)
                        size += len(chunk)
                        output.write(chunk)
                manifest["files"].append({"source": str(source), "name": name, "kind": kind,
                                          "size": size, "sha256": digest.hexdigest()})
            for record in manifest["files"]:
                if image_digest(record["source"]) != record["sha256"]:
                    raise OSError("バックアップ中に元データが変更されました。処理終了後に再試行してください。")
            archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False))
        os.replace(temporary, destination)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return {"path": str(destination), "files": len(sources), "missing": manifest["missing"]}


def restore_bundle(source, parent):
    parent = writable_directory(parent)
    with zipfile.ZipFile(source) as archive:
        if len(archive.infolist()) > 30000 or archive.getinfo("manifest.json").file_size > 32 * 1024 * 1024:
            raise ValueError("バックアップの項目数または管理情報が大きすぎます。")
        manifest = json.loads(archive.read("manifest.json"))
        if manifest.get("format") != "ReceiptLedgerData" or manifest.get("version") != 1:
            raise ValueError("この形式の一括バックアップには対応していません。")
        try:
            validate_session(manifest.get("session"))
        except (TypeError, KeyError, AttributeError) as exc:
            raise ValueError("バックアップ内の下書きを解釈できません。") from exc
        records = manifest["files"]
        if not isinstance(records, list) or len({r["name"] for r in records}) != len(records):
            raise ValueError("バックアップのファイル一覧が不正です。")
        if len({r["source"].casefold() for r in records}) != len(records):
            raise ValueError("バックアップ内に重複した保存先があります。")
        for record in records:
            if (not re.fullmatch(r"files/\d{6}\.(xlsx|png|jpg|jpeg|heic|heif|webp)", record["name"])
                    or not re.fullmatch(r"[0-9a-f]{64}", record["sha256"])
                    or record["kind"] not in ("excel", "backup", "image")
                    or (record["kind"] in ("excel", "backup") and not record["name"].endswith(".xlsx"))
                    or (record["kind"] == "image" and record["name"].endswith(".xlsx"))
                    or not isinstance(record["size"], int) or record["size"] < 0
                    or archive.getinfo(record["name"]).file_size != record["size"]):
                raise ValueError("バックアップ内のファイル情報が不正です。")
        if sum(record["size"] for record in records) > 20 * 1024 ** 3:
            raise ValueError("20GBを超えるバックアップは復元できません。")
        allowed = {record["name"] for record in records} | {"manifest.json"}
        if set(archive.namelist()) != allowed or len(archive.namelist()) != len(allowed):
            raise ValueError("バックアップに未定義のファイルが含まれています。")
        # Stage in a private temporary directory; no existing user data is overwritten.
        final = parent / ("ReceiptLedger-restored-" + datetime.now().strftime("%Y%m%d-%H%M%S-%f"))
        with tempfile.TemporaryDirectory(prefix=".restore-", dir=parent) as staging:
            staging = Path(staging)
            mapping = {}
            destinations = {}
            for record in records:
                suffix = Path(record["name"]).suffix
                if record["kind"] == "image":
                    # Identical pictures imported under different paths are separate queue entries.
                    relative = Path("images") / (record["sha256"] + "-" + Path(record["name"]).stem + suffix)
                else:
                    # Preserve workbook names where possible; suffix protects same-name ledgers.
                    filename = Path(record["source"]).name
                    if not filename or any(c in filename for c in '<>:"/\\|?*'):
                        filename = Path(record["name"]).name
                    relative = Path("excel") / filename
                    if relative.as_posix().casefold() in destinations:
                        relative = Path("excel") / (Path(filename).stem + "-" + Path(record["name"]).stem + suffix)
                destinations[relative.as_posix().casefold()] = True
                target = staging / relative
                target.parent.mkdir(parents=True, exist_ok=True)
                digest = hashlib.sha256()
                with archive.open(record["name"]) as input_file, target.open("wb") as output:
                    while chunk := input_file.read(1024 * 1024):
                        digest.update(chunk)
                        output.write(chunk)
                if digest.hexdigest() != record["sha256"]:
                    raise ValueError("バックアップの検証に失敗しました。データが破損しています。")
                mapping[record["source"].casefold()] = str(final / relative)
            def remap(path):
                return mapping.get(path.casefold(), path)
            # Rebase image references inside each workbook without regenerating its UI sheets.
            from openpyxl import load_workbook
            for record in records:
                if record["kind"] not in ("excel", "backup"):
                    continue
                relative = Path(mapping[record["source"].casefold()]).relative_to(final)
                ledger = ExcelLedger(staging / relative)
                ledger.load()  # Validate the application workbook schema first.
                book = load_workbook(ledger.path)
                try:
                    for row in book["支出一覧"].iter_rows(min_row=2):
                        value = row[8].value
                        if not value:
                            continue
                        original = Path(value)
                        if not original.is_absolute():
                            original = (Path(record["source"]).parent / original).resolve()
                        rebased = remap(str(original))
                        try:
                            rebased = os.path.relpath(rebased, (final / relative).parent)
                        except ValueError:
                            pass
                        row[8].value, row[8].data_type = rebased, "s"
                    book.save(ledger.path)
                finally:
                    book.close()
            settings = Settings(**{k: v for k, v in manifest["settings"].items() if k in Settings.__dataclass_fields__})
            original_workbook = str(Path(settings.workbook_path).resolve())
            settings.workbook_path = mapping.get(original_workbook.casefold(), str(final / "excel" / "レシート家計簿.xlsx"))
            settings.image_folder = str(final / "images")
            session = map_session_paths(manifest["session"], remap)
            session["workbook"] = settings.workbook_path
            (staging / "images").mkdir(exist_ok=True)
            (staging / "excel").mkdir(exist_ok=True)
            write_json_atomic(staging / "restored-settings.json", asdict(settings))
            write_json_atomic(staging / "restored-session.json", session)
            os.replace(staging, final)
    return {"folder": str(final), "settings": asdict(settings), "session": session,
            "missing": manifest.get("missing", [])}

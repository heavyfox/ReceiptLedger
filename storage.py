"""Local file storage, backup retention and integrity-checked image copies."""
import os
import re
import shutil
import tempfile
from pathlib import Path

from core import image_digest


def writable_directory(folder):
    folder = Path(folder).expanduser().resolve()
    folder.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryFile(dir=folder):
        pass
    return folder


def copy_image(source, folder):
    source = Path(source).resolve()
    folder = writable_directory(folder)
    digest = image_digest(source)
    target = folder / (digest + source.suffix.lower())
    if target.exists() and image_digest(target) == digest:
        return target, digest
    fd, name = tempfile.mkstemp(prefix=".image-", dir=folder)
    os.close(fd)
    temporary = Path(name)
    try:
        shutil.copy2(source, temporary)
        if image_digest(temporary) != digest:
            raise OSError("コピー中に画像が変更されました。再試行してください。")
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target, digest


def regular_backups(workbook):
    path = Path(workbook)
    pattern = re.compile(re.escape(path.stem) + r"\.backup-\d{8}-\d{6}-\d{6}\.xlsx\Z")
    if not path.parent.exists():
        return []
    return sorted((item for item in path.parent.iterdir() if item.is_file() and not item.is_symlink()
                   and pattern.fullmatch(item.name)), key=lambda item: item.name, reverse=True)


def prune_backups(workbook, keep):
    if not isinstance(keep, int) or not 1 <= keep <= 999:
        raise ValueError("バックアップの保持数は1〜999を指定してください。")
    removed = []
    for path in regular_backups(workbook)[keep:]:
        path.unlink()
        removed.append(path)
    return removed


def file_usage(paths):
    seen, size = set(), 0
    for path in paths:
        path = Path(path)
        if path.is_file() and not path.is_symlink():
            key = str(path.resolve()).casefold()
            if key not in seen:
                seen.add(key)
                size += path.stat().st_size
    return len(seen), size


def folder_files(folder):
    # Do not follow linked directories outside the selected storage folder.
    folder = Path(folder)
    if not folder.exists():
        return
    for root, dirs, names in os.walk(folder, followlinks=False):
        dirs[:] = [name for name in dirs if not (Path(root) / name).is_symlink()
                   and not os.path.isjunction(Path(root) / name)]
        for name in names:
            path = Path(root) / name
            if not path.is_symlink():
                yield path


def size_text(size):
    if size < 1024:
        return f"{size:,} B"
    if size < 1024 ** 2:
        return f"{size / 1024:.1f} KB"
    return f"{size / (1024 ** 3):.2f} GB" if size >= 1024 ** 3 else f"{size / (1024 ** 2):.1f} MB"


def copy_workbook_rebased(source, destination):
    from core import ExcelLedger
    from openpyxl import load_workbook
    source, destination = Path(source), Path(destination)
    records = {receipt.receipt_id: receipt for receipt in ExcelLedger(source).load()}
    folder = writable_directory(destination.parent)
    fd, name = tempfile.mkstemp(prefix=".restore-excel-", suffix=".xlsx", dir=folder)
    os.close(fd)
    book = load_workbook(source)
    try:
        for row in book["支出一覧"].iter_rows(min_row=2):
            receipt = records.get(str(row[0].value))
            if receipt and receipt.image_path:
                try:
                    value = os.path.relpath(receipt.image_path, folder)
                except ValueError:
                    value = receipt.image_path
                row[8].value, row[8].data_type = value, "s"
        book.save(name)
        os.replace(name, destination)
    finally:
        book.close()
        Path(name).unlink(missing_ok=True)

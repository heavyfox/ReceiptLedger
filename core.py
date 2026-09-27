"""Excel-backed receipt ledger and validation logic."""

from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.chart.data_source import AxDataSource, StrData, StrVal
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink


CATEGORIES = ("食費", "日用品", "交通", "医療", "娯楽", "住居", "光熱費", "その他")
EXPENSE_SHEET = "支出一覧"
ITEM_SHEET = "明細"
SUMMARY_SHEET = "月別集計"
DASHBOARD_SHEET = "ダッシュボード"
RECEIPT_SHEET = "レシート別"
EXPENSE_COLUMNS = (
    "レシートID", "購入日", "店舗", "費目", "金額", "支払方法", "メモ",
    "画像ハッシュ", "画像パス", "税額", "値引き額", "作成日時", "更新日時",
)
ITEM_COLUMNS = ("レシートID", "行番号", "品目", "数量", "単価", "行金額", "費目", "補足")
SUMMARY_COLUMNS = ("年月", "費目", "支出合計", "件数")
YEN_FORMAT = '#,##0"円";[Red]-#,##0"円"'
NAVY = "17324A"
BLUE = "17698A"
PALE = "EAF3F8"
LIGHT = "F5F8FA"
MUTED = "647D8F"


def _display_cell(ws, row: int, column: int, value, *, bold=False, color=NAVY,
                  fill=None, size=11, align="left"):
    cell = ws.cell(row, column)
    if isinstance(value, str):
        _safe_text(cell, value)
    else:
        cell.value = value
    cell.font = Font(name="Yu Gothic", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def _title_row(ws, row: int, text: str):
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=2, max_col=8):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    _display_cell(ws, row, 2, text, bold=True, color="FFFFFF", fill=NAVY, size=17)
    ws.row_dimensions[row].height = 38


def _build_receipt_sheet(ws, receipts: list[Receipt]) -> dict[str, int]:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BLUE
    for col, width in {"A": 3, "B": 32, "C": 12, "D": 12, "E": 12,
                       "F": 15, "G": 16, "H": 42}.items():
        ws.column_dimensions[col].width = width
    _title_row(ws, 2, "レシート別 ｜ 購入した商品の内訳")
    _display_cell(ws, 3, 2, "レシートごとに商品、個数、単価、金額、補足を表示します。アプリで保存すると更新されます。", color=MUTED)
    ws.merge_cells("B3:H3")
    ws.freeze_panes = "B5"
    anchors = {}
    row = 5
    ordered = sorted(receipts, key=lambda r: (r.purchased_on, r.updated_at, r.receipt_id), reverse=True)
    if not ordered:
        _display_cell(ws, row, 2, "レシートはまだ登録されていません。", color=MUTED)
        return anchors
    for number, receipt in enumerate(ordered, 1):
        anchors[receipt.receipt_id] = row
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        heading = f"{number:02d}   {receipt.purchased_on:%Y/%m/%d}   {receipt.merchant}    税込合計 {receipt.amount:,} 円"
        _display_cell(ws, row, 2, heading, bold=True, color="FFFFFF", fill=BLUE, size=13)
        for col in range(3, 9):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.row_dimensions[row].height = 32
        row += 1
        tax_text = f"{receipt.tax:,} 円" if receipt.tax is not None else "未記入"
        discount_text = f"{receipt.discount:,} 円" if receipt.discount is not None else "未記入"
        details = f"費目: {receipt.category}    支払方法: {receipt.payment_method or '未記入'}    うち消費税額: {tax_text}    値引き額: {discount_text}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        _display_cell(ws, row, 2, details, fill=PALE, size=10)
        ws.row_dimensions[row].height = 26
        row += 1
        if receipt.memo:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            _display_cell(ws, row, 2, f"メモ: {receipt.memo}", fill=LIGHT, size=10)
            ws.row_dimensions[row].height = 30
            row += 1
        labels = {"B": "商品名", "E": "個数", "F": "単価", "G": "金額", "H": "詳細・補足"}
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        for col in range(2, 9):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE)
        for col, label in labels.items():
            _display_cell(ws, row, ord(col) - 64, label, bold=True, fill=PALE, size=10)
        ws.row_dimensions[row].height = 25
        row += 1
        if not receipt.items:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            _display_cell(ws, row, 2, "商品明細は未登録です。", color=MUTED)
            ws.row_dimensions[row].height = 25
            row += 1
        for index, item in enumerate(receipt.items):
            band = "FFFFFF" if index % 2 == 0 else LIGHT
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            for col in range(2, 9):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=band)
            _display_cell(ws, row, 2, item.name, fill=band)
            _display_cell(ws, row, 5, item.quantity, fill=band, align="right")
            _display_cell(ws, row, 6, item.unit_price, fill=band, align="right")
            _display_cell(ws, row, 7, item.amount, fill=band, align="right")
            _display_cell(ws, row, 8, item.note, fill=band)
            for col in (6, 7):
                ws.cell(row, col).number_format = YEN_FORMAT
            ws.row_dimensions[row].height = 35 if item.note else 26
            row += 1
        value = sum(item.amount for item in receipt.items if item.amount is not None)
        treatment = tax_treatment(receipt)
        tax_label, tax_note = {
            "added": ("消費税（小計に加算）", "商品小計との差と一致"),
            "included": ("うち消費税", "商品明細の金額に含まれています"),
            "unknown": ("消費税（内訳・参考）", "元画像で税の扱いを確認してください"),
        }[treatment]
        summary_rows = (
            ("商品小計（税抜）" if treatment == "added" else "商品明細の合計",
             value if receipt.items else None, ""),
            (tax_label, receipt.tax, "未記入" if receipt.tax is None else tax_note),
            ("合計（税込）", receipt.amount, ""),
        )
        difference = discrepancy(receipt)
        for label, amount, note in summary_rows:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
            _display_cell(ws, row, 2, label, bold=True, fill=PALE)
            _display_cell(ws, row, 7, amount, bold=True, fill=PALE, align="right").number_format = YEN_FORMAT
            if label == "合計（税込）" and difference:
                note = f"要確認: 明細との差 {difference:,} 円"
            _display_cell(ws, row, 8, note, color=MUTED, fill=PALE, size=10)
            ws.row_dimensions[row].height = 27
            row += 1
        row += 1
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:3"
    return anchors


def _build_dashboard(ws, receipts: list[Receipt], anchors: dict[str, int]):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "3E9A87"
    for col, width in {"A": 3, "B": 21, "C": 17, "D": 17, "E": 4,
                       "F": 20, "G": 19, "H": 22}.items():
        ws.column_dimensions[col].width = width
    _title_row(ws, 2, "レシート家計簿 ｜ ダッシュボード")
    ws.merge_cells("B3:H3")
    _display_cell(ws, 3, 2, "支出の全体像を確認できます。商品ごとの内容は「レシート別」シートへ。", color=MUTED)
    today_month = date.today().strftime("%Y-%m")
    cards = (("B", "C", "累計支出", sum(r.amount for r in receipts), YEN_FORMAT),
             ("D", "F", "今月の支出", sum(r.amount for r in receipts if r.purchased_on.strftime("%Y-%m") == today_month), YEN_FORMAT),
             ("G", "H", "レシート件数", len(receipts), '#,##0"件"'))
    for start, end, label, value, fmt in cards:
        c1, c2 = ord(start)-64, ord(end)-64
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c2)
        _display_cell(ws, 5, c1, label, bold=True, color=MUTED, fill=PALE, size=10)
        cell = _display_cell(ws, 6, c1, value, bold=True, color=BLUE, fill=PALE, size=20)
        cell.number_format = fmt
        for row in (5, 6):
            for col in range(c1+1, c2+1):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE)
    ws.row_dimensions[5].height = 23
    ws.row_dimensions[6].height = 38
    _display_cell(ws, 9, 2, "月別支出", bold=True, color="FFFFFF", fill=NAVY)
    _display_cell(ws, 9, 3, "金額", bold=True, color="FFFFFF", fill=NAVY)
    _display_cell(ws, 9, 4, "件数", bold=True, color="FFFFFF", fill=NAVY)
    _display_cell(ws, 9, 6, "費目別支出", bold=True, color="FFFFFF", fill=NAVY)
    _display_cell(ws, 9, 7, "金額", bold=True, color="FFFFFF", fill=NAVY)
    months: dict[str, list[int]] = {}
    categories: dict[str, int] = {}
    for receipt in receipts:
        month = receipt.purchased_on.strftime("%Y-%m")
        monthly = months.setdefault(month, [0, 0])
        monthly[0] += receipt.amount
        monthly[1] += 1
        categories[receipt.category] = categories.get(receipt.category, 0) + receipt.amount
    for index, (month, (amount, count)) in enumerate(sorted(months.items(), reverse=True), 10):
        fill = LIGHT if index % 2 == 0 else "FFFFFF"
        _display_cell(ws, index, 2, month, fill=fill)
        _display_cell(ws, index, 3, amount, fill=fill, align="right").number_format = YEN_FORMAT
        _display_cell(ws, index, 4, count, fill=fill, align="right")
    for index, (category, amount) in enumerate(sorted(categories.items(), key=lambda x: -x[1]), 10):
        fill = LIGHT if index % 2 == 0 else "FFFFFF"
        _display_cell(ws, index, 6, category, fill=fill)
        _display_cell(ws, index, 7, amount, fill=fill, align="right").number_format = YEN_FORMAT
    recent_row = max(12, 11 + len(months), 11 + len(categories)) + 2
    ws.merge_cells(start_row=recent_row, start_column=2, end_row=recent_row, end_column=8)
    _display_cell(ws, recent_row, 2, "最近のレシート  |  店舗名を押すと明細へ移動", bold=True, color="FFFFFF", fill=NAVY)
    recent_row += 1
    for col, label in ((2, "購入日"), (3, "店舗"), (6, "費目"), (7, "金額"), (8, "商品数")):
        _display_cell(ws, recent_row, col, label, bold=True, fill=PALE)
    for index, receipt in enumerate(sorted(receipts, key=lambda r: (r.purchased_on, r.updated_at), reverse=True)[:12], recent_row + 1):
        fill = LIGHT if index % 2 == 0 else "FFFFFF"
        _display_cell(ws, index, 2, receipt.purchased_on.strftime("%Y/%m/%d"), fill=fill)
        ws.merge_cells(start_row=index, start_column=3, end_row=index, end_column=4)
        link = _display_cell(ws, index, 3, receipt.merchant, color=BLUE, fill=fill)
        link.hyperlink = Hyperlink(ref=link.coordinate,
                                   location=f"'{RECEIPT_SHEET}'!B{anchors[receipt.receipt_id]}",
                                   display=receipt.merchant)
        link.font = Font(name="Yu Gothic", size=11, color=BLUE, underline="single")
        _display_cell(ws, index, 6, receipt.category, fill=fill)
        _display_cell(ws, index, 7, receipt.amount, fill=fill, align="right").number_format = YEN_FORMAT
        _display_cell(ws, index, 8, len(receipt.items), fill=fill, align="right")
        ws.row_dimensions[index].height = 25
    ws.freeze_panes = "B9"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1


def _build_monthly_summary(ws, receipts):
    """A chronological monthly overview, refreshed whenever the ledger is saved."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.tabColor = "3E9A87"
    ws.column_dimensions["A"].width = 3
    for col in range(2, 13):
        ws.column_dimensions[get_column_letter(col)].width = 16 if col < 4 else 13
    ws.column_dimensions["D"].width = 9
    ws.merge_cells("B2:L2")
    _display_cell(ws, 2, 2, "月別集計 ｜ 支出の推移と内訳", bold=True, color="FFFFFF", fill=NAVY, size=20)
    ws.row_dimensions[2].height = 42
    ws.merge_cells("B3:L3")
    _display_cell(ws, 3, 2, "金額は税込・費目はレシート単位。— は0円です。アプリで「Excelを保存」すると更新します。", color=MUTED, size=10)
    ws.row_dimensions[3].height = 28
    months = {}
    for receipt in receipts:
        key = receipt.purchased_on.strftime("%Y-%m")
        group = months.setdefault(key, {"count": 0, "total": 0, **dict.fromkeys(CATEGORIES, 0)})
        group["count"] += 1
        group["total"] += receipt.amount
        group[receipt.category] += receipt.amount
    cards = ((2, 4, "累計支出（税込）", sum(r.amount for r in receipts), YEN_FORMAT),
             (5, 8, "レシート件数", len(receipts), '#,##0"件"'),
             (9, 12, "記録のある月", len(months), '0"か月"'))
    for first, last, label, value, fmt in cards:
        for row in (5, 6):
            ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
            for col in range(first, last + 1):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE)
        _display_cell(ws, 5, first, label, bold=True, color=MUTED, fill=PALE, size=10)
        _display_cell(ws, 6, first, value, bold=True, color=BLUE, fill=PALE, size=22).number_format = fmt
    ws.row_dimensions[5].height = 25
    ws.row_dimensions[6].height = 40
    ws.merge_cells("B8:D8")
    ws.merge_cells("E8:L8")
    _display_cell(ws, 8, 2, "月ごとの支出", bold=True, color=BLUE, size=12)
    _display_cell(ws, 8, 5, "費目別の内訳", bold=True, color=BLUE, size=12)
    ws.row_dimensions[8].height = 30
    for col, label in enumerate(("年月", "税込合計", "件数", *CATEGORIES), 2):
        _display_cell(ws, 9, col, label, bold=True, color="FFFFFF", fill=NAVY if col < 5 else BLUE, align="center")
    ws.row_dimensions[9].height = 30
    keys = []
    if months:
        first, last = min(months), max(months)
        year, month = map(int, first.split("-"))
        while f"{year:04d}-{month:02d}" <= last:
            keys.append(f"{year:04d}-{month:02d}")
            year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    for row, key in enumerate(keys, 10):
        group = months.get(key, {"count": 0, "total": 0, **dict.fromkeys(CATEGORIES, 0)})
        fill = LIGHT if row % 2 == 0 else "FFFFFF"
        year, month = map(int, key.split("-"))
        _display_cell(ws, row, 2, date(year, month, 1), fill=fill).number_format = 'yyyy"年"m"月"'
        for col, value in enumerate((group["total"], group["count"], *(group[c] for c in CATEGORIES)), 3):
            cell = _display_cell(ws, row, col, value, fill=PALE if col == 3 else fill,
                                 bold=col == 3, color=MUTED if value == 0 else NAVY, align="center" if col == 4 else "right")
            cell.number_format = '#,##0"件"' if col == 4 else '#,##0"円";[Red]-#,##0"円";"—"'
        ws.row_dimensions[row].height = 35
    if keys:
        last_row = 9 + len(keys)
        ws.auto_filter.ref = f"B9:L{last_row}"
        ws.conditional_formatting.add(f"C10:C{last_row}", DataBarRule(start_type="min", end_type="max", color="D5E8EB", showValue=True))
        total_row = last_row + 1
        _display_cell(ws, total_row, 2, "合計", bold=True, fill=PALE)
        for col in range(3, 13):
            value = sum(ws.cell(row, col).value for row in range(10, last_row + 1))
            _display_cell(ws, total_row, col, value, bold=True, fill=PALE, align="right").number_format = '#,##0"件"' if col == 4 else YEN_FORMAT
        ws.row_dimensions[total_row].height = 34
        years = sorted({key[:4] for key in keys})
        for offset, year in enumerate(years):
            year_keys = [key for key in keys if key[:4] == year]
            first_row = 10 + keys.index(year_keys[0])
            end_row = first_row + len(year_keys) - 1
            chart = BarChart()
            chart.title = f"{year}年 ｜ 月ごとの支出（税込）"
            chart.y_axis.title = "金額（円）"
            chart.add_data(Reference(ws, min_col=3, min_row=first_row, max_row=end_row))
            chart.series[0].cat = AxDataSource(strLit=StrData(ptCount=len(year_keys),
                pt=[StrVal(idx=i, v=f"{int(key[5:])}月") for i, key in enumerate(year_keys)]))
            chart.legend = None
            chart.width, chart.height = 30, 12
            chart.style = 10
            chart.gapWidth = 65
            chart.series[0].graphicalProperties.solidFill = BLUE
            chart.series[0].graphicalProperties.line.noFill = True
            chart.y_axis.numFmt = '#,##0"円"'
            # Excel requires explicit visible axes; other preview engines can
            # display labels even when these settings are omitted.
            chart.x_axis.delete = False
            chart.y_axis.delete = False
            chart.x_axis.axPos = "b"
            chart.y_axis.axPos = "l"
            chart.y_axis.tickLblPos = "nextTo"
            chart.x_axis.crosses = "autoZero"
            chart.y_axis.crosses = "autoZero"
            chart.x_axis.tickLblPos = "low"
            chart.x_axis.tickLblSkip = 1
            chart.x_axis.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1200, b=True)))])
            chart.dLbls = DataLabelList(showVal=True, showLegendKey=False,
                showCatName=False, showSerName=False, dLblPos="outEnd", numFmt='#,##0"円"')
            chart.dLbls.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1100, b=True)))])
            values = [ws.cell(row, 3).value for row in range(first_row, end_row + 1)]
            chart.y_axis.scaling.max = max(1, max(values)) * 1.2
            chart.y_axis.scaling.min = min(0, min(values) * 1.2)
            ws.add_chart(chart, f"B{total_row + 3 + offset * 27}")
        bottom = total_row + 3 + len(years) * 27
    else:
        ws.merge_cells("B10:L10")
        _display_cell(ws, 10, 2, "レシートを記録すると、ここに月ごとの支出とグラフが表示されます。", color=MUTED)
        ws.row_dimensions[10].height = 35
        bottom = 11
    ws.freeze_panes = "A10"
    ws.print_area = f"B2:L{bottom}"
    ws.print_title_rows = "9:9"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


class LedgerError(Exception):
    pass


class WorkbookChanged(LedgerError):
    pass


@dataclass
class LineItem:
    name: str = ""
    quantity: float | None = None
    unit_price: int | None = None
    amount: int | None = None
    category: str = ""
    note: str = ""


@dataclass
class Receipt:
    receipt_id: str = field(default_factory=lambda: uuid4().hex)
    purchased_on: date = field(default_factory=date.today)
    merchant: str = ""
    amount: int = 0
    category: str = "その他"
    payment_method: str = ""
    memo: str = ""
    image_hash: str = ""
    image_path: str = ""
    tax: int | None = None
    discount: int | None = None
    created_at: datetime = field(default_factory=datetime.now)
    updated_at: datetime = field(default_factory=datetime.now)
    items: list[LineItem] = field(default_factory=list)


def image_digest(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_receipt(receipt: Receipt) -> None:
    if not receipt.merchant.strip():
        raise LedgerError("店舗名を入力してください。")
    if receipt.purchased_on > date.today():
        raise LedgerError("購入日が未来の日付です。")
    if receipt.category not in CATEGORIES:
        raise LedgerError("費目を選択してください。")
    if abs(receipt.amount) > 1_000_000_000:
        raise LedgerError("金額が大きすぎます。")


def tax_treatment(receipt: Receipt) -> str:
    """Reconcile only complete item amounts against the printed total and tax."""
    if not receipt.items or any(item.amount is None for item in receipt.items):
        return "unknown"
    subtotal = sum(item.amount for item in receipt.items)
    if subtotal == receipt.amount:
        return "included"
    if receipt.tax is not None and subtotal + receipt.tax == receipt.amount:
        return "added"
    return "unknown"


def discrepancy(receipt: Receipt) -> int | None:
    """Flag unexplained differences; the printed tax can explain external tax."""
    amounts = [item.amount for item in receipt.items if item.amount is not None]
    if not amounts:
        return None
    if tax_treatment(receipt) == "added":
        return 0
    return receipt.amount - sum(amounts)


def duplicate_candidates(receipt: Receipt, existing: list[Receipt]) -> list[Receipt]:
    matches = []
    for other in existing:
        if other.receipt_id == receipt.receipt_id:
            continue
        same_image = bool(receipt.image_hash and receipt.image_hash == other.image_hash)
        same_fields = (
            receipt.purchased_on == other.purchased_on
            and receipt.amount == other.amount
            and receipt.merchant.strip().casefold() == other.merchant.strip().casefold()
        )
        if same_image or same_fields:
            matches.append(other)
    return matches


def _fingerprint(path: Path) -> str | None:
    return image_digest(path) if path.exists() else None


def _cell_text(value) -> str:
    return "" if value is None else str(value)


def _safe_text(cell, value: str) -> None:
    cell.value = value
    cell.data_type = "s"  # Prevent Excel from interpreting receipt text as a formula.


def _to_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _to_datetime(value) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    return datetime.fromisoformat(str(value))


class ExcelLedger:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.fingerprint: str | None = None
        self.backup_limit: int | None = None
        self.maintenance_warning = ""

    def delete(self) -> Path:
        """Delete the loaded workbook after retaining a recoverable copy."""
        try:
            if not self.path.is_file():
                raise LedgerError("削除する Excel ファイルがありません。")
            if self.fingerprint is None:
                raise LedgerError("家計簿を正常に読み込んでから削除してください。")
            if _fingerprint(self.path) != self.fingerprint:
                raise WorkbookChanged("Excel ファイルが外部で変更されました。再読み込みして内容を確認してください。")
            stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
            backup = self.path.with_name(f"{self.path.stem}.before-delete-{stamp}.xlsx")
            shutil.copy2(self.path, backup)
            if _fingerprint(self.path) != self.fingerprint:
                raise WorkbookChanged("削除の準備中に Excel ファイルが変更されました。再読み込みしてください。")
            self.path.unlink()
            self.fingerprint = None
            return backup
        except LedgerError:
            raise
        except OSError as exc:
            raise LedgerError("Excel ファイルを削除できません。Excel で開いている場合は閉じ、フォルダーへの書き込み権限を確認してください。") from exc

    def load(self) -> list[Receipt]:
        if not self.path.exists():
            self.fingerprint = None
            return []
        try:
            book = load_workbook(self.path, data_only=True)
            if EXPENSE_SHEET not in book or ITEM_SHEET not in book:
                raise LedgerError("このファイルはアプリの家計簿形式ではありません。")
            expenses = book[EXPENSE_SHEET]
            items_sheet = book[ITEM_SHEET]
            if tuple(c.value for c in expenses[1])[: len(EXPENSE_COLUMNS)] != EXPENSE_COLUMNS:
                raise LedgerError("支出一覧の列が変更されています。バックアップを確認してください。")
            if tuple(c.value for c in items_sheet[1])[: len(ITEM_COLUMNS)] != ITEM_COLUMNS:
                raise LedgerError("明細の列が変更されています。バックアップを確認してください。")
            receipts: dict[str, Receipt] = {}
            for values in expenses.iter_rows(min_row=2, values_only=True):
                if not values[0]:
                    continue
                row = list(values) + [None] * max(0, len(EXPENSE_COLUMNS) - len(values))
                if row[1] is None or row[4] is None:
                    raise LedgerError("購入日または金額が空欄の行があります。Excel ファイルを確認してください。")
                receipt = Receipt(
                    receipt_id=_cell_text(row[0]), purchased_on=_to_date(row[1]),
                    merchant=_cell_text(row[2]), category=_cell_text(row[3]) or "その他",
                    amount=int(row[4] or 0), payment_method=_cell_text(row[5]),
                    memo=_cell_text(row[6]), image_hash=_cell_text(row[7]),
                    image_path=_cell_text(row[8]), tax=int(row[9]) if row[9] is not None else None,
                    discount=int(row[10]) if row[10] is not None else None,
                    created_at=_to_datetime(row[11]) if row[11] else datetime.now(),
                    updated_at=_to_datetime(row[12]) if row[12] else datetime.now(),
                )
                if receipt.receipt_id in receipts:
                    raise LedgerError("レシート ID が重複しています。")
                if receipt.image_path and not Path(receipt.image_path).is_absolute():
                    receipt.image_path = str((self.path.parent / receipt.image_path).resolve())
                receipts[receipt.receipt_id] = receipt
            for values in items_sheet.iter_rows(min_row=2, values_only=True):
                if not values[0]:
                    continue
                row = list(values) + [None] * max(0, len(ITEM_COLUMNS) - len(values))
                receipt = receipts.get(_cell_text(row[0]))
                if receipt is None:
                    raise LedgerError("明細に対応する支出がありません。")
                receipt.items.append(LineItem(
                    name=_cell_text(row[2]), quantity=float(row[3]) if row[3] is not None else None,
                    unit_price=int(row[4]) if row[4] is not None else None,
                    amount=int(row[5]) if row[5] is not None else None,
                    category=_cell_text(row[6]), note=_cell_text(row[7]),
                ))
            self.fingerprint = _fingerprint(self.path)
            return list(receipts.values())
        except LedgerError:
            raise
        except Exception as exc:
            raise LedgerError(f"Excel ファイルを読み込めません: {exc}") from exc

    def save(self, receipts: list[Receipt]) -> Path | None:
        if _fingerprint(self.path) != self.fingerprint:
            raise WorkbookChanged("Excel ファイルが外部で変更されました。再読み込みしてください。")
        if len({r.receipt_id for r in receipts}) != len(receipts):
            raise LedgerError("レシート ID が重複しています。")
        for receipt in receipts:
            validate_receipt(receipt)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        backup = None
        temp_path = None
        try:
            book = load_workbook(self.path) if self.path.exists() else Workbook()
            if book.active.title == "Sheet" and book.active.max_row == 1 and book.active["A1"].value is None:
                book.remove(book.active)
            for name in (DASHBOARD_SHEET, RECEIPT_SHEET, EXPENSE_SHEET, ITEM_SHEET, SUMMARY_SHEET):
                if name in book:
                    book.remove(book[name])
            dashboard_ws = book.create_sheet(DASHBOARD_SHEET, 0)
            receipt_ws = book.create_sheet(RECEIPT_SHEET, 1)
            expense_ws = book.create_sheet(EXPENSE_SHEET, 2)
            item_ws = book.create_sheet(ITEM_SHEET, 3)
            summary_ws = book.create_sheet(SUMMARY_SHEET, 4)
            for ws, columns in ((expense_ws, EXPENSE_COLUMNS), (item_ws, ITEM_COLUMNS)):
                ws.append(columns)
                ws.freeze_panes = "A2"
                ws.sheet_view.showGridLines = False
                ws.row_dimensions[1].height = 25
                for cell in ws[1]:
                    cell.font = Font(name="Yu Gothic", bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor=NAVY)
                    cell.alignment = Alignment(vertical="center")
            expense_ws.sheet_properties.tabColor = "5A91AC"
            item_ws.sheet_properties.tabColor = "73AFC0"
            summary_ws.sheet_properties.tabColor = "8CB2A4"
            for receipt in sorted(receipts, key=lambda r: (r.purchased_on, r.receipt_id)):
                image_path = receipt.image_path
                if image_path and Path(image_path).is_absolute():
                    try:
                        image_path = os.path.relpath(image_path, self.path.parent.resolve())
                    except ValueError:
                        pass  # Different Windows drives require an absolute reference.
                expense_ws.append((
                    receipt.receipt_id, receipt.purchased_on, receipt.merchant,
                    receipt.category, receipt.amount, receipt.payment_method,
                    receipt.memo, receipt.image_hash, image_path,
                    receipt.tax, receipt.discount, receipt.created_at, receipt.updated_at,
                ))
                row_no = expense_ws.max_row
                for index in (1, 3, 4, 6, 7, 8, 9):
                    _safe_text(expense_ws.cell(row_no, index), str(expense_ws.cell(row_no, index).value or ""))
                expense_ws.cell(row_no, 2).number_format = "yyyy/mm/dd"
                expense_ws.cell(row_no, 5).number_format = YEN_FORMAT
                for col in (10, 11):
                    expense_ws.cell(row_no, col).number_format = YEN_FORMAT
                for col in (12, 13):
                    expense_ws.cell(row_no, col).number_format = "yyyy/mm/dd hh:mm"
                expense_ws.row_dimensions[row_no].height = 29
                expense_fill = LIGHT if row_no % 2 == 0 else "FFFFFF"
                for cell in expense_ws[row_no]:
                    cell.fill = PatternFill("solid", fgColor=expense_fill)
                    cell.font = Font(name="Yu Gothic", size=10, color=NAVY)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                for index, item in enumerate(receipt.items, 1):
                    item_ws.append((receipt.receipt_id, index, item.name, item.quantity,
                                    item.unit_price, item.amount, item.category, item.note))
                    item_row = item_ws.max_row
                    for col in (1, 3, 7, 8):
                        _safe_text(item_ws.cell(item_row, col), str(item_ws.cell(item_row, col).value or ""))
                    item_fill = PALE if row_no % 2 == 0 else "FFFFFF"
                    for cell in item_ws[item_row]:
                        cell.fill = PatternFill("solid", fgColor=item_fill)
                        cell.font = Font(name="Yu Gothic", size=10, color=NAVY)
                        cell.alignment = Alignment(vertical="center", wrap_text=True)
                    item_ws.row_dimensions[item_row].height = 29
                    item_ws.cell(item_row, 4).number_format = "General"
                    for col in (5, 6):
                        item_ws.cell(item_row, col).number_format = YEN_FORMAT
            _build_monthly_summary(summary_ws, receipts)
            for ws, widths in ((expense_ws, [34, 16, 26, 15, 15, 18, 34, 20, 38, 13, 13, 21, 21]),
                               (item_ws, [34, 10, 32, 12, 14, 14, 15, 32])):
                for index, width in enumerate(widths, 1):
                    ws.column_dimensions[get_column_letter(index)].width = width
                ws.auto_filter.ref = ws.dimensions
            anchors = _build_receipt_sheet(receipt_ws, receipts)
            _build_dashboard(dashboard_ws, receipts, anchors)
            for row_no in range(2, expense_ws.max_row + 1):
                receipt_id = expense_ws.cell(row_no, 1).value
                if receipt_id in anchors:
                    cell = expense_ws.cell(row_no, 3)
                    cell.hyperlink = Hyperlink(ref=cell.coordinate,
                                               location=f"'{RECEIPT_SHEET}'!B{anchors[receipt_id]}",
                                               display=str(cell.value))
                    cell.font = Font(name="Yu Gothic", size=10, color=BLUE, underline="single")
            book.active = 0
            fd, name = tempfile.mkstemp(prefix=".receipt-ledger-", suffix=".xlsx", dir=self.path.parent)
            os.close(fd)
            temp_path = Path(name)
            book.save(temp_path)
            if self.path.exists():
                stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
                backup = self.path.with_name(f"{self.path.stem}.backup-{stamp}.xlsx")
                shutil.copy2(self.path, backup)
            os.replace(temp_path, self.path)
            temp_path = None
            self.fingerprint = _fingerprint(self.path)
            self.maintenance_warning = ""
            if self.backup_limit is not None:
                try:
                    from storage import prune_backups
                    prune_backups(self.path, self.backup_limit)
                except (OSError, ValueError) as exc:
                    # The ledger is already committed; never report a failed save here.
                    self.maintenance_warning = f"Excelは保存済みですが、古いバックアップを整理できません: {exc}"
            return backup
        except LedgerError:
            raise
        except PermissionError as exc:
            raise LedgerError("Excel ファイルに保存できません。Excel で開いている場合は閉じて再試行してください。") from exc
        except Exception as exc:
            raise LedgerError(f"Excel ファイルに保存できません: {exc}") from exc
        finally:
            if temp_path and temp_path.exists():
                temp_path.unlink(missing_ok=True)

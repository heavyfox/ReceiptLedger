import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from core import ExcelLedger, LineItem, Receipt, discrepancy, tax_treatment


class TaxTotalsTests(unittest.TestCase):
    def receipt(self, amounts, total, tax):
        return Receipt(purchased_on=date(2026, 9, 2), merchant="テスト店", amount=total, tax=tax,
                       items=[LineItem(f"商品{i}", 1, amount, amount) for i, amount in enumerate(amounts)])

    def test_external_tax_is_shown_below_items_and_explains_difference(self):
        receipt = self.receipt([139, 149, 149, 209], 697, 51)
        self.assertEqual(tax_treatment(receipt), "added")
        self.assertEqual(discrepancy(receipt), 0)
        with tempfile.TemporaryDirectory() as temp:
            ledger = ExcelLedger(Path(temp) / "tax.xlsx")
            ledger.save([receipt])
            book = load_workbook(ledger.path, data_only=True)
            rows = {row[1]: (row[6], row[7]) for row in book["レシート別"].iter_rows(values_only=True) if row[1]}
            self.assertEqual(rows["商品小計（税抜）"][0], 646)
            self.assertEqual(rows["消費税（小計に加算）"][0], 51)
            self.assertEqual(rows["合計（税込）"], (697, None))
            self.assertEqual(book["支出一覧"]["E2"].value, 697)
            self.assertEqual(book["支出一覧"]["J2"].value, 51)
            self.assertEqual(book["月別集計"]["C10"].value, 697)
            self.assertEqual(len(ledger.load()[0].items), 4)
            book.close()

    def test_included_tax_is_not_added_twice(self):
        receipt = self.receipt([300, 397], 697, 51)
        self.assertEqual(tax_treatment(receipt), "included")
        self.assertEqual(discrepancy(receipt), 0)

    def test_unknown_or_incomplete_items_keep_difference_warning(self):
        for amounts, total, tax, expected in (([600], 697, 51, 97),
                                             ([646], 697, None, 51),
                                             ([646, None], 697, 51, 51)):
            receipt = self.receipt(amounts, total, tax)
            self.assertEqual(tax_treatment(receipt), "unknown")
            self.assertEqual(discrepancy(receipt), expected)

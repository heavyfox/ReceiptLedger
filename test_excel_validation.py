from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from core import CATEGORIES, ExcelLedger, LedgerError, LineItem, Receipt
from excel_validation import ExportValidationError, NS, validate_export


class ExcelOutputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root / "ledger.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def export(self, receipts):
        ledger = ExcelLedger(self.path)
        ledger.load()
        ledger.save(receipts)
        validate_export(self.path, receipts, CATEGORIES)
        return ledger

    def mutate(self, part, edit):
        changed = self.root / "changed.xlsx"
        with zipfile.ZipFile(self.path) as source, zipfile.ZipFile(changed, "w") as output:
            for info in source.infolist():
                raw = source.read(info.filename)
                if info.filename == part:
                    tree = ET.fromstring(raw)
                    edit(tree)
                    raw = ET.tostring(tree)
                output.writestr(info, raw)
        return changed

    def test_twelve_months_tax_zero_missing_tax_and_refund(self):
        receipts = [Receipt(purchased_on=date(2025, month, 15), merchant="架空店", category="食費",
                            amount=month * 100, tax=0 if month == 1 else None) for month in range(1, 13)]
        receipts += [Receipt(purchased_on=date(2025, 2, 20), merchant="外税店", category="日用品", amount=110,
                             tax=10, items=[LineItem("品物", 1, 100, 100, "日用品")]),
                     Receipt(purchased_on=date(2025, 3, 1), merchant="返金", category="食費", amount=-50, tax=-4)]
        self.export(receipts)
        book = load_workbook(self.path, data_only=True)
        try:
            summary = book["月別集計"]
            self.assertEqual(summary["C11"].value, 310)  # gross, not 320
            self.assertEqual(summary["C12"].value, 250)
            self.assertEqual(summary["C22"].value, 7860)
            recorded = {row[0]: row[9] for row in book["支出一覧"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(recorded[receipts[0].receipt_id], 0)
            self.assertIsNone(recorded[receipts[1].receipt_id])
            self.assertEqual(recorded[receipts[-1].receipt_id], -4)
        finally:
            book.close()

    def test_year_boundary_gaps_and_large_amount(self):
        receipts = [Receipt(purchased_on=date(2024, 12, 1), merchant="A", category="食費", amount=0),
                    Receipt(purchased_on=date(2026, 2, 1), merchant="B", category="日用品", amount=1_000_000_000, tax=100)]
        self.export(receipts)
        book = load_workbook(self.path, data_only=True)
        try:
            self.assertEqual(book["月別集計"]["C11"].value, 0)
            self.assertEqual(book["月別集計"]["C24"].value, 1_000_000_000)
            self.assertEqual(len(book["月別集計"]._charts), 3)
        finally:
            book.close()

    def test_empty_and_single_month_workbooks(self):
        self.export([])
        self.export([Receipt(purchased_on=date(2025, 1, 1), merchant="A", category="食費", amount=10)])

    def test_missing_month_axis_labels_or_amount_labels_are_rejected(self):
        receipts = [Receipt(purchased_on=date(2025, 1, 1), merchant="A", category="食費", amount=10)]
        self.export(receipts)
        mutations = [lambda tree: tree.find(".//c:catAx/c:delete", NS).set("val", "1"),
                     lambda tree: tree.find(".//c:cat/c:strLit/c:pt/c:v", NS).__setattr__("text", ""),
                     lambda tree: tree.find(".//c:dLbls/c:showVal", NS).set("val", "0"),
                     lambda tree: tree.find(".//c:val/c:numRef/c:f", NS).__setattr__("text", "'月別集計'!$C$11")]
        for edit in mutations:
            with self.subTest(edit=edit):
                changed = self.mutate("xl/charts/chart1.xml", edit)
                with self.assertRaises(ExportValidationError):
                    validate_export(changed, receipts, CATEGORIES)

    def test_tax_missing_from_saved_xml_is_rejected(self):
        receipts = [Receipt(purchased_on=date(2025, 1, 1), merchant="A", category="食費", amount=110, tax=10)]
        self.export(receipts)
        def erase_tax(tree):
            cell = tree.find('.//s:c[@r="J2"]', NS)
            cell.remove(cell.find("s:v", NS))
        changed = self.mutate("xl/worksheets/sheet3.xml", erase_tax)
        with self.assertRaisesRegex(ExportValidationError, "消費税"):
            validate_export(changed, receipts, CATEGORIES)

    def test_failed_output_validation_keeps_previous_file_and_cleans_temp(self):
        receipt = Receipt(purchased_on=date(2025, 1, 1), merchant="A", category="食費", amount=110)
        ledger = self.export([receipt])
        previous = self.path.read_bytes()
        with patch("core.validate_export", side_effect=ExportValidationError("グラフの軸が非表示です。")):
            with self.assertRaises(LedgerError):
                ledger.save([receipt])
        self.assertEqual(self.path.read_bytes(), previous)
        self.assertFalse(list(self.root.glob(".receipt-ledger-*.xlsx")))
        self.assertFalse(list(self.root.glob("*.backup-*.xlsx")))

    def test_existing_1904_dates_and_user_sheet_are_preserved(self):
        receipt = Receipt(purchased_on=date(2025, 1, 1), merchant="A", category="食費", amount=110)
        self.export([receipt])
        book = load_workbook(self.path)
        book.epoch = CALENDAR_MAC_1904
        book.create_sheet("個人メモ")["A1"] = "残す内容"
        book.save(self.path)
        book.close()
        self.export([receipt])
        book = load_workbook(self.path)
        try:
            self.assertEqual(book["個人メモ"]["A1"].value, "残す内容")
            self.assertEqual(book.epoch, CALENDAR_MAC_1904)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()

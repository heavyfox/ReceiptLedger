import os
import tempfile
import threading
import unittest
from collections import Counter
from datetime import date
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PIL import Image
from openpyxl import load_workbook
from PySide6.QtCore import Qt, QElapsedTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox, QStyleOptionViewItem

from app import MainWindow
from batch import ReceiptDraft
from config import Settings
from core import ExcelLedger, LedgerError
from lm_client import _parse_receipt


def result_for(path):
    return _parse_receipt({
        "receipt_date": date.today().isoformat(), "merchant": Path(path).stem,
        "total_yen": 300, "category": "食費",
        "items": [{"name": Path(path).stem + "の商品", "quantity": 2,
                   "unit_price_yen": 150, "line_total_yen": 300, "note": "産地情報"}],
    })


class BatchTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.env = patch.dict(os.environ, {"RECEIPT_LEDGER_DATA_DIR": str(self.root)})
        self.env.start()
        Settings(workbook_path=str(self.root / "家計簿.xlsx"), model="local-test").save()
        with patch("app.get_token", return_value=""):
            self.window = MainWindow()
        self.release = threading.Event()
        self.warning = patch("app.QMessageBox.warning").start()
        self.info = patch("app.QMessageBox.information").start()

    def tearDown(self):
        self.release.set()
        if self.window.batch_worker:
            self.window.batch_worker.requestInterruption()
            self.window.batch_worker.wait(5000)
        self.application.processEvents()
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            self.window.close()
        patch.stopall()
        self.env.stop()
        self.temp.cleanup()

    def wait_until(self, predicate):
        timer = QElapsedTimer()
        timer.start()
        while not predicate() and timer.elapsed() < 5000:
            self.application.processEvents()
            QTest.qWait(5)
        self.assertTrue(predicate(), "Background reading did not finish in time")

    def add_images(self, names):
        paths = []
        for index, name in enumerate(names):
            path = self.root / name
            if path.suffix == ".heic":
                path.write_bytes((Path(__file__).resolve().parent / "tests/fixtures/receipt.heic").read_bytes())
            else:
                Image.new("RGB", (200, 300), (index * 60, 30, 90)).save(path, format="PNG")
            paths.append(str(path))
        self.window._enqueue_images(paths)
        return paths

    def mark_all(self):
        for i in range(self.window.queue.count()):
            self.window.queue.setCurrentRow(i)
            self.window._mark_reviewed()

    def test_failure_retry_corrections_and_one_batch_save(self):
        paths = self.add_images(["first.png", "bad.png", "last.heic"])
        calls = []
        attempts = Counter()

        def extract(path, model):
            calls.append(path)
            attempts[path] += 1
            if path == paths[1] and attempts[path] == 1:
                raise RuntimeError("一時的な接続エラー")
            return result_for(path)

        with patch("app.LMStudioClient") as client:
            client.return_value.extract.side_effect = extract
            self.window._read_batch()
            self.wait_until(lambda: not self.window.workers)
            self.assertEqual([self.window.entries[p].state for p in paths], ["ready", "failed", "ready"])
            self.assertFalse(self.window.ledger.path.exists())
            self.window.queue.setCurrentRow(0)
            self.window.merchant_edit.setText("修正した店")
            self.window.item_table.item(0, 4).setText("補足も修正")
            self.window.queue.setCurrentRow(2)
            self.window.queue.setCurrentRow(0)
            self.assertEqual(self.window.merchant_edit.text(), "修正した店")
            self.assertEqual(self.window.item_table.item(0, 4).text(), "補足も修正")
            self.window._retry_failed()
            self.wait_until(lambda: not self.window.workers)
        self.assertEqual(calls, [*paths, paths[1]])
        self.mark_all()
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes), patch.object(self.window.ledger, "save", wraps=self.window.ledger.save) as save:
            self.window._save_batch()
            save.assert_called_once()
            self.window._save_batch()
            save.assert_called_once()
        records = ExcelLedger(self.window.ledger.path).load()
        self.assertEqual(len(records), 3)
        records = {r.merchant: r for r in records}
        self.assertEqual(records["修正した店"].items[0].name, "firstの商品")
        self.assertEqual(records["修正した店"].items[0].note, "補足も修正")
        self.assertEqual(records["last"].items[0].name, "lastの商品")
        self.assertTrue(all(self.window.entries[p].state == "saved" for p in paths))
        self.warning.assert_not_called()

    def test_stop_keeps_current_result_and_resume_only_unread(self):
        paths = self.add_images(["one.png", "two.png", "three.png"])
        started = threading.Event()
        calls = []

        def extract(path, model):
            calls.append(path)
            if len(calls) == 1:
                started.set()
                self.release.wait(5)
            return result_for(path)

        with patch("app.LMStudioClient") as client:
            client.return_value.extract.side_effect = extract
            self.window._read_batch()
            self.wait_until(started.is_set)
            # Changing the selection while a request runs must not mix results.
            self.window.queue.setCurrentRow(1)
            self.window._stop_batch()
            self.release.set()
            self.wait_until(lambda: not self.window.workers)
            self.assertEqual(calls, [paths[0]])
            self.assertEqual([self.window.entries[p].state for p in paths], ["ready", "pending", "pending"])
            self.assertEqual(self.window.merchant_edit.text(), "")
            self.window._read_batch()
            self.wait_until(lambda: not self.window.workers)
        self.assertEqual(calls, paths)
        self.assertEqual(self.window.merchant_edit.text(), "two")
        self.assertEqual(self.window.batch_progress.value(), 2)

    def test_save_failure_preserves_reviewed_drafts_and_can_retry(self):
        paths = self.add_images(["one.png", "two.png"])
        for path in paths:
            self.window._batch_image_succeeded(path, result_for(path))
        self.mark_all()
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes), patch.object(self.window.ledger, "save", side_effect=LedgerError("Excelで開いています")):
            self.window._save_batch()
        self.assertEqual(self.window.receipts, [])
        self.assertTrue(all(self.window.entries[p].reviewed for p in paths))
        self.assertTrue(all(self.window.entries[p].checked for p in paths))
        self.assertTrue(all(self.window.entries[p].state == "ready" for p in paths))
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            self.window._save_batch()
        self.assertEqual(len(ExcelLedger(self.window.ledger.path).load()), 2)

    def test_invalid_checked_draft_blocks_entire_batch_after_edit(self):
        paths = self.add_images(["one.png", "two.png"])
        for path in paths:
            self.window._batch_image_succeeded(path, result_for(path))
        self.mark_all()
        self.window.queue.setCurrentRow(0)
        self.window.date_edit.setText("not a date")
        self.window.queue.setCurrentRow(1)
        self.assertFalse(self.window.entries[paths[0]].reviewed)
        self.assertTrue(self.window.entries[paths[0]].checked)
        self.window.queue.item(0).setCheckState(Qt.CheckState.Checked)
        with patch.object(self.window.ledger, "save") as save:
            self.window._save_batch()
            save.assert_not_called()
        self.assertFalse(self.window.entries[paths[0]].reviewed)
        self.warning.assert_called_once()

    def test_combined_button_saves_checked_only_and_cancel_keeps_drafts(self):
        paths = self.add_images(["one.png", "two.png", "unread.png"])
        for path in paths[:2]:
            self.window._batch_image_succeeded(path, result_for(path))
        self.window.queue.item(0).setCheckState(Qt.CheckState.Checked)
        self.assertTrue(self.window.entries[paths[0]].checked)
        self.assertFalse(self.window.entries[paths[0]].reviewed)
        self.assertFalse(self.window.queue.item(2).flags() & Qt.ItemFlag.ItemIsUserCheckable)
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.No), patch.object(self.window.ledger, "save") as save:
            self.window.batch_save_btn.click()
            save.assert_not_called()
        self.assertTrue(self.window.entries[paths[0]].checked)
        self.assertEqual(self.window.entries[paths[0]].state, "ready")
        # An independently confirmed receipt without a check must be excluded.
        self.window.queue.setCurrentRow(1)
        self.window._mark_reviewed()
        self.window.queue.item(1).setCheckState(Qt.CheckState.Unchecked)
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes), patch.object(self.window.ledger, "save", wraps=self.window.ledger.save) as save:
            self.window.batch_save_btn.click()
            self.window._save_batch()
            save.assert_called_once()
        self.assertEqual(len(ExcelLedger(self.window.ledger.path).load()), 1)
        self.assertEqual(self.window.entries[paths[0]].state, "saved")
        self.assertFalse(self.window.entries[paths[0]].checked)
        self.assertEqual(self.window.entries[paths[1]].state, "ready")
        self.assertEqual(self.window.entries[paths[2]].state, "pending")

    def test_edit_and_bulk_review_completed_receipt_while_next_is_reading(self):
        paths = self.add_images(["one.png", "two.png"])
        second_started = threading.Event()

        def extract(path, model):
            if path == paths[1]:
                second_started.set()
                self.release.wait(5)
            return result_for(path)

        with patch("app.LMStudioClient") as client:
            client.return_value.extract.side_effect = extract
            self.window._read_batch()
            self.wait_until(lambda: second_started.is_set() and self.window.entries[paths[0]].state == "ready")
            self.assertTrue(self.window.receipt_editor.isEnabled())
            self.assertFalse(self.window.save_btn.isEnabled())
            self.assertFalse(self.window.batch_save_btn.isEnabled())
            self.window.merchant_edit.selectAll()
            QTest.keyClicks(self.window.merchant_edit, "edited during reading")
            self.window.item_table.item(0, 4).setText("処理中に修正した補足")
            self.window.queue.item(0).setCheckState(Qt.CheckState.Checked)
            self.assertFalse(self.window.batch_save_btn.isEnabled())
            self.window.theme_selector.setCurrentIndex(self.window.theme_selector.findData("dark"))
            self.assertEqual(self.window.settings.theme, "dark")
            self.assertTrue(self.window.entries[paths[0]].checked)
            self.window.queue.setCurrentRow(1)
            self.assertFalse(self.window.receipt_editor.isEnabled())
            self.window.queue.setCurrentRow(0)
            self.assertTrue(self.window.receipt_editor.isEnabled())
            self.assertEqual(self.window.merchant_edit.text(), "edited during reading")
            self.release.set()
            self.wait_until(lambda: not self.window.workers)
        self.assertEqual(self.window.merchant_edit.text(), "edited during reading")
        self.assertEqual(self.window.item_table.item(0, 4).text(), "処理中に修正した補足")
        self.assertTrue(self.window.entries[paths[0]].checked)
        self.assertFalse(self.window.entries[paths[1]].reviewed)
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            self.window._save_batch()
        saved = ExcelLedger(self.window.ledger.path).load()
        self.assertEqual(len(saved), 1)
        self.assertEqual(saved[0].merchant, "edited during reading")
        self.assertEqual(saved[0].items[0].note, "処理中に修正した補足")
        self.warning.assert_not_called()

    def test_close_during_batch_finishes_current_then_saves_and_stops(self):
        paths = self.add_images(["one.png", "two.png", "three.png"])
        started = threading.Event()
        calls = []
        def extract(path, model):
            calls.append(path)
            if path == paths[1]:
                started.set()
                self.release.wait(5)
            return result_for(path)
        self.window.show()
        with patch("app.LMStudioClient") as client:
            client.return_value.extract.side_effect = extract
            self.window._read_batch()
            self.wait_until(lambda: started.is_set() and self.window.entries[paths[0]].state == "ready")
            self.window.memo_edit.setText("終了前の修正")
            self.window.close()
            self.assertTrue(self.window._close_after_batch)
            self.release.set()
            self.wait_until(lambda: not self.window.workers and not self.window.isVisible())
        saved = self.window.session_store.load()
        self.assertEqual(calls, paths[:2])
        self.assertEqual([entry["state"] for entry in saved["entries"]], ["ready", "ready", "pending"])
        self.assertEqual(saved["entries"][0]["draft"]["memo"], "終了前の修正")

    def test_large_checkbox_mouse_and_keyboard_in_both_themes(self):
        paths = self.add_images(["ready.png", "pending.png"])
        self.window._batch_image_succeeded(paths[0], result_for(paths[0]))
        self.window.show()
        self.application.processEvents()
        for theme in ("light", "dark"):
            self.window.theme_box.setCurrentIndex(self.window.theme_box.findData(theme))
            self.application.processEvents()
            option = QStyleOptionViewItem()
            option.rect = self.window.queue.visualItemRect(self.window.queue.item(0))
            # Click near the large indicator's edge, outside the old native checkbox.
            pos = self.window.queue_delegate.indicator_rect(option).bottomRight()
            pos.setX(pos.x() - 2)
            pos.setY(pos.y() - 2)
            QTest.mouseClick(self.window.queue.viewport(), Qt.MouseButton.LeftButton, pos=pos)
            self.assertTrue(self.window.entries[paths[0]].checked)
            self.assertFalse(self.window.entries[paths[0]].reviewed)
            self.window.queue.setCurrentRow(0)
            QTest.keyClick(self.window.queue, Qt.Key.Key_Space)
            self.assertFalse(self.window.entries[paths[0]].checked)
            option.rect = self.window.queue.visualItemRect(self.window.queue.item(1))
            QTest.mouseClick(self.window.queue.viewport(), Qt.MouseButton.LeftButton,
                             pos=self.window.queue_delegate.indicator_rect(option).center())
            self.assertFalse(self.window.entries[paths[1]].checked)

    def test_theme_persists_and_keeps_unsaved_form(self):
        self.window.merchant_edit.setText("編集中の店舗")
        self.window.memo_edit.setText("保存前のメモ")
        self.window.theme_selector.setCurrentIndex(self.window.theme_selector.findData("dark"))
        self.assertEqual(Settings.load().theme, "dark")
        self.assertEqual(self.window.merchant_edit.text(), "編集中の店舗")
        self.assertEqual(self.window.memo_edit.text(), "保存前のメモ")
        with patch("app.get_token", return_value=""):
            reopened = MainWindow()
        try:
            self.assertEqual(reopened.settings.theme, "dark")
            self.assertEqual(reopened.theme_box.currentData(), "dark")
        finally:
            reopened.close()
        # Existing installations can load their settings without a theme field.
        (self.root / "settings.json").write_text('{"model":"local-test"}', encoding="utf-8")
        self.assertEqual(Settings.load().theme, "system")

    def test_system_theme_follows_windows_and_manual_override_wins(self):
        self.assertEqual(self.window.settings.theme, "system")
        for current in ("light", "dark"):
            with patch("app.system_theme", return_value=current):
                self.window._sync_system_theme()
                self.assertEqual(self.window.effective_theme, current)
        self.window.theme_box.setCurrentIndex(self.window.theme_box.findData("light"))
        with patch("app.system_theme", return_value="dark"):
            self.window._sync_system_theme()
            self.assertEqual(self.window.effective_theme, "light")
            self.assertEqual(self.window.theme_selector.currentData(), "light")
            self.window.theme_selector.setCurrentIndex(self.window.theme_selector.findData("system"))
            self.assertEqual(self.window.effective_theme, "dark")
            self.assertEqual(self.window.theme_box.currentData(), "system")
        self.assertEqual(Settings.load().theme, "system")

    def test_extracted_gross_total_and_tax_saved_without_double_counting(self):
        cases = [(1080, 80, [], 80), (2180, None, [80, 100], 180),
                 (500, 0, [], 0), (800, None, [None, 50], None)]
        paths = self.add_images([f"tax-{i}.png" for i in range(len(cases))])
        for path, (total, tax, components, expected_tax) in zip(paths, cases):
            extracted = _parse_receipt({"receipt_date": date.today().isoformat(),
                "merchant": Path(path).stem, "total_yen": total, "tax_yen": tax,
                "tax_components_yen": components, "category": "食費"})
            self.window._batch_image_succeeded(path, extracted)
            self.window.queue.setCurrentRow(paths.index(path))
            self.assertEqual(self.window.amount_edit.text(), str(total))
            self.assertEqual(self.window.tax_edit.text(), "" if expected_tax is None else str(expected_tax))
            self.window.queue.item(paths.index(path)).setCheckState(Qt.CheckState.Checked)
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            self.window.batch_save_btn.click()
        saved = ExcelLedger(self.window.ledger.path).load()
        self.assertEqual({r.merchant: (r.amount, r.tax) for r in saved},
                         {f"tax-{i}": (c[0], c[3]) for i, c in enumerate(cases)})
        book = load_workbook(self.window.ledger.path, data_only=True)
        self.assertEqual({row[2]: (row[4], row[9]) for row in book["支出一覧"].iter_rows(min_row=2, values_only=True)},
                         {f"tax-{i}": (c[0], c[3]) for i, c in enumerate(cases)})
        self.assertEqual(book["月別集計"]["C2"].value, 4560)
        displayed = "\n".join(str(c.value) for row in book["レシート別"] for c in row if c.value is not None)
        self.assertIn("税込合計 2,180 円", displayed)
        self.assertIn("うち消費税額: 180 円", displayed)
        book.close()

    def test_scheduled_reread_locks_old_draft_until_stop_or_result(self):
        paths = self.add_images(["one.png", "two.png"])
        for path in paths:
            self.window._batch_image_succeeded(path, result_for(path))
        first_started = threading.Event()

        def extract(path, model):
            first_started.set()
            self.release.wait(5)
            return result_for(path)

        with patch("app.LMStudioClient") as client:
            client.return_value.extract.side_effect = extract
            self.window._begin_batch(paths)
            self.wait_until(first_started.is_set)
            self.window.queue.setCurrentRow(1)
            self.assertFalse(self.window.receipt_editor.isEnabled())
            self.assertFalse(self.window._can_review(self.window.entries[paths[1]]))
            self.assertFalse(self.window.queue.item(1).flags() & Qt.ItemFlag.ItemIsUserCheckable)
            self.window._stop_batch()
            self.release.set()
            self.wait_until(lambda: not self.window.workers)
        self.assertTrue(self.window.receipt_editor.isEnabled())
        self.assertTrue(self.window._can_review(self.window.entries[paths[1]]))
        self.assertEqual(self.window.merchant_edit.text(), "two")

    def test_batch_confirmation_includes_duplicates_and_amount_difference(self):
        paths = self.add_images(["one.png", "two.png"])
        for path in paths:
            self.window._batch_image_succeeded(path, result_for(path))
        self.window.receipts = [self.window.entries[paths[0]].draft.to_receipt(paths[0])]
        self.window.entries[paths[1]].draft.amount = "500"
        self.mark_all()
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.No) as question:
            self.window._save_batch()
        message = question.call_args.args[2]
        self.assertIn("重複候補", message)
        self.assertIn("+200", message)
        self.assertFalse(self.window.ledger.path.exists())


if __name__ == "__main__":
    unittest.main()

import json
import os
import tempfile
import threading
import unittest
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook

from core import ExcelLedger, LineItem, Receipt, WorkbookChanged, discrepancy, duplicate_candidates
from lm_client import LMStudioClient, LMStudioError, normalize_base_url, _parse_receipt


class LedgerTests(unittest.TestCase):
    def test_receipt_views_show_items_and_details(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "household.xlsx"
            store = ExcelLedger(path)
            first = Receipt(
                purchased_on=date(2026, 9, 20), merchant="青果店", amount=450,
                category="食費", payment_method="現金", tax=30, memo="夕食用",
                items=[LineItem("りんご", 2, 150, 300, "食費", "青森産"),
                       LineItem("にんじん", 1, 150, 150, "食費", "3本入り")],
            )
            second = Receipt(
                purchased_on=date(2026, 9, 21), merchant="薬局", amount=200,
                category="日用品", items=[LineItem("洗剤", 1, 200, 200, "日用品", "詰め替え")],
            )
            store.save([first, second])
            book = load_workbook(path)
            self.assertEqual(book.sheetnames[:5], ["ダッシュボード", "レシート別", "支出一覧", "明細", "月別集計"])
            self.assertEqual(book.active.title, "ダッシュボード")
            detail = book["レシート別"]
            visible = [str(cell.value) for row in detail for cell in row if cell.value is not None]
            for expected in ("りんご", "青森産", "にんじん", "3本入り", "洗剤", "詰め替え"):
                self.assertIn(expected, visible)
            self.assertEqual(book["明細"]["H2"].value, "青森産")
            self.assertEqual(book["明細"]["D2"].value, 2)
            self.assertEqual(book["明細"]["E2"].value, 150)
            self.assertEqual(book["支出一覧"]["C2"].hyperlink.location, "'レシート別'!B13")
            self.assertEqual([r.items[0].note for r in store.load()], ["青森産", "詰め替え"])

    def test_form_preserves_item_note_and_calculates_line_amount(self):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication, QTableWidgetItem
        from app import MainWindow

        with tempfile.TemporaryDirectory() as temp, patch.dict(os.environ, {"RECEIPT_LEDGER_DATA_DIR": temp}):
            application = QApplication.instance() or QApplication([])
            window = MainWindow()
            try:
                window.date_edit.setText(date.today().isoformat())
                window.merchant_edit.setText("青果店")
                window.amount_edit.setText("450")
                window.item_table.insertRow(0)
                for col, value in enumerate(("りんご", "2", "150", "", "青森産")):
                    window.item_table.setItem(0, col, QTableWidgetItem(value))
                receipt = window._collect_form()
                self.assertEqual((receipt.items[0].quantity, receipt.items[0].unit_price,
                                  receipt.items[0].amount, receipt.items[0].note), (2, 150, 300, "青森産"))
                window.receipts = [receipt]
                window._refresh_ledger()
                window.ledger_table.selectRow(0)
                window._edit_selected()
                self.assertEqual(window.item_table.item(0, 4).text(), "青森産")
            finally:
                window.close()

    def test_monthly_totals_include_refunds_once(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "household.xlsx"
            store = ExcelLedger(path)
            records = [
                Receipt(purchased_on=date(2026, 9, 1), merchant="店A", amount=1000, category="食費",
                        items=[LineItem("品物", 1, 1000, 1000)]),
                Receipt(purchased_on=date(2026, 9, 2), merchant="店A", amount=-200, category="食費"),
                Receipt(purchased_on=date(2026, 8, 1), merchant="店B", amount=500, category="日用品"),
            ]
            store.save(records)
            book = load_workbook(path, data_only=True)
            values = {(row[0], row[1]): (row[2], row[3]) for row in
                      book["月別集計"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(values[("2026-09", "食費")], (800, 2))
            self.assertEqual(values[("2026-08", "日用品")], (500, 1))
            self.assertEqual(sum(r.amount for r in store.load()), 1300)

    def test_round_trip_summary_formula_text_and_backup(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "household.xlsx"
            store = ExcelLedger(path)
            receipt = Receipt(
                purchased_on=date(2026, 9, 20), merchant="=HYPERLINK(\"bad\")",
                amount=680, category="食費", image_hash="abc",
                items=[LineItem("牛乳", 1, 180, 180), LineItem("パン", 2, 250, 500)],
            )
            self.assertIsNone(store.save([receipt]))
            loaded = store.load()
            self.assertEqual(loaded[0].merchant, receipt.merchant)
            self.assertEqual(loaded[0].items[1].amount, 500)
            book = load_workbook(path)
            self.assertEqual(book["支出一覧"]["C2"].data_type, "s")
            self.assertEqual(book["月別集計"]["C2"].value, 680)
            receipt.amount = 700
            backup = store.save([receipt])
            self.assertTrue(backup.exists())
            self.assertEqual(load_workbook(backup)["支出一覧"]["E2"].value, 680)

    def test_external_change_and_duplicate(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "household.xlsx"
            store = ExcelLedger(path)
            first = Receipt(purchased_on=date(2026, 9, 20), merchant="店A", amount=100, image_hash="same")
            store.save([first])
            second = Receipt(purchased_on=date(2026, 9, 20), merchant="店A", amount=100, image_hash="same")
            self.assertEqual(len(duplicate_candidates(second, [first])), 1)
            self.assertIsNone(discrepancy(first))
            book = load_workbook(path)
            book["支出一覧"]["G2"] = "外部編集"
            book.save(path)
            with self.assertRaises(WorkbookChanged):
                store.save([first])


class _LMHandler(BaseHTTPRequestHandler):
    calls = []

    def log_message(self, *_):
        pass

    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps({"data": [{"id": "qwen/qwen3.8-27b"}]}).encode())

    def do_POST(self):
        data = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        self.calls.append(data)
        extracted = {
            "receipt_date": "2026-09-20", "merchant": "店A", "total_yen": 300,
            "tax_yen": 27, "discount_yen": None, "payment_method": "現金",
            "category": "食費", "items": [{"name": "パン", "quantity": 1,
                                         "unit_price_yen": 300, "line_total_yen": 300}],
            "uncertain_fields": [],
        }
        answer = {"choices": [{"message": {"content": json.dumps(extracted)}}]}
        body = json.dumps(answer).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


class LMClientTests(unittest.TestCase):
    def test_null_arrays_and_formatted_yen_from_local_model(self):
        result = _parse_receipt({"total_yen": "￥２，５９５円", "tax_yen": "192",
                                 "tax_components_yen": None, "items": None, "uncertain_fields": None})
        self.assertEqual((result.total_yen, result.tax_yen, result.items), (2595, 192, []))
        with self.assertRaises(ValueError):
            _parse_receipt({"total_yen": 123.5})

    def test_model_wrappers_and_truncation_retry(self):
        response = lambda text, reason: {"choices": [{"finish_reason": reason, "message": {"content": text}}]}
        client = LMStudioClient("http://localhost:1234/v1")
        requests = []
        responses = [response('{"total_yen": 2595}', "length"),
                     response('<think>reasoning</think>\n```json\n{"total_yen":2595,"tax_yen":192,"tax_components_yen":null}\n```', "stop")]
        def request(method, path, body):
            requests.append(body["max_tokens"])
            return responses.pop(0)
        with patch("lm_client._image_data_url", return_value="data:image/jpeg;base64,test"), patch.object(client, "_request", side_effect=request):
            result = client.extract("sample.heic", "local-model")
        self.assertEqual((result.total_yen, result.tax_yen), (2595, 192))
        self.assertEqual(requests, [3000, 6000])

    def test_incomplete_response_never_becomes_a_success(self):
        client = LMStudioClient("http://localhost:1234/v1")
        response = {"choices": [{"finish_reason": "length", "message": {"content": '{"total_yen": 2595'}}]}
        with patch("lm_client._image_data_url", return_value="data:image/jpeg;base64,test"), patch.object(client, "_request", return_value=response) as request:
            with self.assertRaisesRegex(LMStudioError, "上限.*"):
                client.extract("sample.heic", "local-model")
            self.assertEqual(request.call_count, 2)

    def test_tax_total_is_not_added_to_breakdown_or_gross_total(self):
        extracted = _parse_receipt({"total_yen": 2180, "tax_yen": 180,
                                    "tax_components_yen": [80, 100]})
        self.assertEqual((extracted.total_yen, extracted.tax_yen), (2180, 180))
        inconsistent = _parse_receipt({"total_yen": 2180, "tax_yen": 190,
                                      "tax_components_yen": [80, 100]})
        self.assertEqual(inconsistent.tax_yen, 190)
        self.assertTrue(any("一致しません" in warning for warning in inconsistent.uncertain_fields))

    def test_reject_public_host(self):
        with self.assertRaises(LMStudioError):
            normalize_base_url("https://example.com:1234/v1")

    def test_models_and_image_extract(self):
        server = ThreadingHTTPServer(("127.0.0.1", 0), _LMHandler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with tempfile.TemporaryDirectory() as temp:
                image_path = Path(temp) / "receipt.png"
                Image.new("RGB", (400, 800), "white").save(image_path)
                client = LMStudioClient(f"http://127.0.0.1:{server.server_port}/v1", "sample")
                self.assertEqual(client.models(), ["qwen/qwen3.8-27b"])
                result = client.extract(image_path, "qwen/qwen3.8-27b")
                self.assertEqual(result.total_yen, 300)
                self.assertEqual(result.tax_yen, 27)
                self.assertTrue(_LMHandler.calls[-1]["messages"][1]["content"][1]["image_url"]["url"].startswith("data:image/jpeg;base64,"))
                self.assertEqual(_LMHandler.calls[-1]["response_format"]["type"], "json_schema")
        finally:
            server.shutdown()
            server.server_close()


if __name__ == "__main__":
    unittest.main()

import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QMessageBox
from openpyxl import load_workbook

from app import MainWindow
from config import Settings
from core import ExcelLedger, LedgerError, Receipt, WorkbookChanged
from lm_client import _parse_receipt


class FileActionsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.env = patch.dict(os.environ, {"RECEIPT_LEDGER_DATA_DIR": str(self.root)})
        self.env.start()
        self.addCleanup(self.env.stop)
        Settings(workbook_path=str(self.root / "家計簿.xlsx")).save()
        with patch("app.get_token", return_value=""):
            self.window = MainWindow()
        self.addCleanup(self.window.close)

    def test_extracted_receipt_is_appended_to_selected_template(self):
        template = self.root / "テンプレート.xlsx"
        ExcelLedger(template).save([])
        initial = load_workbook(template)
        fill = initial["ダッシュボード"]["B2"].fill.fgColor.rgb
        initial.close()
        # Choose a target after reading: the unregistered result must survive.
        self.window._apply_extraction(_parse_receipt({
            "receipt_date": date.today().isoformat(), "merchant": "青果店", "total_yen": 300,
            "category": "食費", "items": [{"name": "りんご", "quantity": 2,
                "unit_price_yen": 150, "line_total_yen": 300, "note": "青森産"}],
        }))
        with patch("app.QFileDialog.getOpenFileName", return_value=(str(template), "Excel")):
            self.window._select_workbook()
        self.assertEqual(self.window.merchant_edit.text(), "青果店")
        with patch("app.QMessageBox.warning") as warning:
            self.window._save_receipt()
            warning.assert_not_called()
        saved = ExcelLedger(template).load()
        self.assertEqual(saved[0].items[0].note, "青森産")
        self.assertEqual(saved[0].items[0].quantity, 2)
        book = load_workbook(template)
        self.assertEqual(book["ダッシュボード"]["B2"].fill.fgColor.rgb, fill)
        self.assertEqual(book["ダッシュボード"]["B6"].value, 300)
        self.assertEqual(book["レシート別"]["B8"].value, "りんご")
        book.close()
        self.assertEqual(Path(Settings.load().workbook_path), template)

    def test_save_creates_empty_workbook_and_save_as_switches_target(self):
        self.window._save_workbook()
        original = self.window.ledger.path
        self.assertTrue(original.exists())
        destination = self.root / "別の家計簿.xlsx"
        with patch("app.QFileDialog.getSaveFileName", return_value=(str(destination), "Excel")):
            self.window._save_workbook_as()
        self.assertEqual(self.window.ledger.path, destination)
        self.window.receipts = [Receipt(merchant="新しい店", amount=123)]
        self.window._save_workbook()
        self.assertEqual(ExcelLedger(original).load(), [])
        self.assertEqual(ExcelLedger(destination).load()[0].amount, 123)
        self.assertEqual(Path(Settings.load().workbook_path), destination)

    def test_delete_cancel_backup_restore_and_new_save(self):
        self.window.receipts = [Receipt(merchant="店", amount=456)]
        self.window._save_workbook()
        path = self.window.ledger.path
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.No):
            self.window._delete_workbook()
        self.assertTrue(path.exists())
        self.assertEqual(len(self.window.receipts), 1)
        with patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes), patch("app.QMessageBox.information"):
            self.window._delete_workbook()
        self.assertFalse(path.exists())
        self.assertEqual(self.window.receipts, [])
        backups = list(self.root.glob("*.before-delete-*.xlsx"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(ExcelLedger(backups[0]).load()[0].amount, 456)
        self.window._save_workbook()
        self.assertEqual(ExcelLedger(path).load(), [])
        with patch("app.QFileDialog.getOpenFileName", return_value=(str(backups[0]), "Excel")), patch("app.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            self.window._restore_backup()
        self.assertEqual(self.window.receipts[0].amount, 456)

    def test_failed_backup_and_external_change_do_not_delete(self):
        ledger = self.window.ledger
        ledger.save([Receipt(merchant="店", amount=200)])
        before = ledger.path.read_bytes()
        with patch("core.shutil.copy2", side_effect=PermissionError("locked")):
            with self.assertRaises(LedgerError):
                ledger.delete()
        self.assertEqual(ledger.path.read_bytes(), before)
        external = load_workbook(ledger.path)
        external["支出一覧"]["G2"] = "外部で編集"
        external.save(ledger.path)
        external.close()
        with self.assertRaises(WorkbookChanged):
            ledger.delete()
        self.assertTrue(ledger.path.exists())

    def test_save_as_existing_workbook_retains_backup(self):
        destination = self.root / "既存.xlsx"
        ExcelLedger(destination).save([Receipt(merchant="保存前", amount=100)])
        self.window.receipts = [Receipt(merchant="保存後", amount=300)]
        with patch("app.QFileDialog.getSaveFileName", return_value=(str(destination), "Excel")):
            self.window._save_workbook_as()
        self.assertEqual(ExcelLedger(destination).load()[0].merchant, "保存後")
        backups = list(self.root.glob("既存.backup-*.xlsx"))
        self.assertEqual(ExcelLedger(backups[0]).load()[0].merchant, "保存前")


if __name__ == "__main__":
    unittest.main()
